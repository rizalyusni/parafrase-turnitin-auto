# -*- coding: utf-8 -*-
"""
Main FastAPI Application for Turnitin Paraphrase Automator.
Features:
- DOCX upload & classification engine
- 100% Strict format lock on tables, equations, citations, and headings
- Database-driven permanent vocabulary substitutions (Padanan Kata)
- Self-validation framework (Tahap 1–4, Checklist A–E)
- Direct download of paraphrased .docx
- Ready for Railway deployment
"""
import os
import shutil
import uuid
from pathlib import Path
from fastapi import FastAPI, UploadFile, File, BackgroundTasks, Depends, HTTPException
from fastapi.responses import HTMLResponse, FileResponse, JSONResponse
from fastapi.middleware.cors import CORSMiddleware
from sqlalchemy.orm import Session
import docx

from config import UPLOAD_DIR, OUTPUT_DIR, PORT, HOST
from database.db import get_db, init_db, SessionLocal
from database.seeds import seed_database
from database.models import PadananKata, IstilahTerkunci, ParaphraseJob, JobParagraph, TypoReport, SystemSetting
from engine.classifier import classify_paragraph, audit_locked_typos
from engine.docx_handler import process_and_save_docx
from engine.paraphraser import batch_paraphrase
from engine.validator import run_document_validation

app = FastAPI(
    title="Turnitin Paraphrase Automator API",
    description="Automated Academic Paraphrasing Tool with Strict Format Locking & SQL Thesaurus",
    version="2.0.0"
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# Startup: initialize database tables and seed initial vocabulary & terms
@app.on_event("startup")
def on_startup():
    init_db()
    db = SessionLocal()
    try:
        seed_database(db)
    finally:
        db.close()
    print("[Server] Database initialized with seed padanan and locked terms.")

# ----------------- UI ROUTE -----------------
@app.get("/", response_class=HTMLResponse)
def index_page():
    html_file = Path(__file__).resolve().parent / "templates" / "index.html"
    if html_file.exists():
        return html_file.read_text(encoding="utf-8")
    return "<h1>Turnitin Paraphrase Automator Running</h1>"

# ----------------- BACKGROUND WORKER -----------------
def execute_paraphrase_job(job_id: str, input_path: str, filename: str):
    db = SessionLocal()
    try:
        job = db.query(ParaphraseJob).filter_by(id=job_id).first()
        if not job:
            return

        job.status = "analyzing"
        job.progress = 20
        db.commit()

        # Step 1: Parse and classify paragraphs
        doc = docx.Document(input_path)
        total_p = len(doc.paragraphs)
        job.total_paragraphs = total_p

        classified_items = []
        narrative_items = []
        locked_indices = set()
        raw_texts = []

        for idx, p in enumerate(doc.paragraphs):
            raw_texts.append(p.text)
            classification, is_locked = classify_paragraph(p, idx)
            classified_items.append((idx, p.style.name, classification, is_locked, p.text))
            if is_locked:
                locked_indices.add(idx)
            else:
                narrative_items.append({"index": idx, "text": p.text.strip()})

        job.narrative_count = len(narrative_items)
        job.locked_count = len(locked_indices)
        job.progress = 40
        job.status = "paraphrasing"
        db.commit()

        # Step 2: Audit typos on locked elements
        typo_findings = audit_locked_typos(raw_texts)
        for tf in typo_findings:
            db.add(TypoReport(
                job_id=job.id,
                paragraph_index=tf["paragraph_index"],
                typo_found=tf["typo_found"],
                saran_perbaikan=tf["saran_perbaikan"],
                catatan=tf["catatan"]
            ))
        db.commit()

        # Step 3: Fetch active padanan kata & locked terms from SQL
        padanan_records = [
            {"kata_asal": p.kata_asal, "padanan_disukai": p.padanan_disukai, "padanan_terlarang": p.padanan_terlarang}
            for p in db.query(PadananKata).filter_by(is_active=True).all()
        ]
        locked_terms = [t.istilah for t in db.query(IstilahTerkunci).filter_by(is_active=True).all()]

        # Step 4: Execute Paraphrasing on narrative items
        paraphrase_results = batch_paraphrase(narrative_items, padanan_records, locked_terms)

        job.progress = 75
        job.status = "formatting"
        db.commit()

        # Step 5: Generate the new .docx with preserved formatting & styles
        base_name = Path(filename).stem
        output_filename = f"{base_name} - Hasil Parafrase Turnitin.docx"
        output_path = str(OUTPUT_DIR / f"{job_id}_{output_filename}")

        process_and_save_docx(input_path, output_path, paraphrase_results)

        # Step 6: Validate output
        job.progress = 90
        job.status = "validating"
        db.commit()

        val_report = run_document_validation(input_path, output_path, locked_indices)

        # Step 7: Record paragraph details in database for audit
        for idx, style_name, classification, is_locked, orig_txt in classified_items:
            paraphrased_txt = paraphrase_results.get(idx, orig_txt)
            db.add(JobParagraph(
                job_id=job.id,
                paragraph_index=idx,
                style_name=style_name,
                classification=classification,
                is_locked=is_locked,
                original_text=orig_txt,
                paraphrased_text=paraphrased_txt,
                audit_notes="Preserved" if is_locked else "Paraphrased with Turnitin format-lock"
            ))

        job.status = "completed"
        job.progress = 100
        job.validation_score = val_report["score"]
        job.result_path = output_path
        db.commit()

    except Exception as e:
        print(f"[Worker Error] {e}")
        job = db.query(ParaphraseJob).filter_by(id=job_id).first()
        if job:
            job.status = "failed"
            job.error_message = str(e)
            db.commit()
    finally:
        db.close()

# ----------------- API ENDPOINTS -----------------
@app.post("/api/upload")
async def upload_document(
    background_tasks: BackgroundTasks,
    file: UploadFile = File(...),
    db: Session = Depends(get_db)
):
    if not file.filename.endswith(".docx"):
        raise HTTPException(status_code=400, detail="Hanya file berekstensi .docx yang didukung.")

    job_id = str(uuid.uuid4())
    save_filename = f"{job_id}_{file.filename}"
    file_path = UPLOAD_DIR / save_filename

    with open(file_path, "wb") as buffer:
        shutil.copyfileobj(file.file, buffer)

    file_size = os.path.getsize(file_path)

    new_job = ParaphraseJob(
        id=job_id,
        filename=file.filename,
        original_size=file_size,
        status="pending",
        progress=5,
        file_path=str(file_path)
    )
    db.add(new_job)
    db.commit()

    background_tasks.add_task(execute_paraphrase_job, job_id, str(file_path), file.filename)

    return {"status": "success", "job_id": job_id, "message": "Dokumen berhasil diunggah dan sedang diproses."}

@app.get("/api/jobs/{job_id}/status")
def get_job_status(job_id: str, db: Session = Depends(get_db)):
    job = db.query(ParaphraseJob).filter_by(id=job_id).first()
    if not job:
        raise HTTPException(status_code=404, detail="Job tidak ditemukan.")

    return {
        "id": job.id,
        "filename": job.filename,
        "status": job.status,
        "progress": job.progress,
        "total_paragraphs": job.total_paragraphs,
        "narrative_count": job.narrative_count,
        "locked_count": job.locked_count,
        "validation_score": job.validation_score,
        "error_message": job.error_message
    }

@app.get("/api/jobs/{job_id}/result")
def download_result(job_id: str, db: Session = Depends(get_db)):
    job = db.query(ParaphraseJob).filter_by(id=job_id).first()
    if not job or not job.result_path or not os.path.exists(job.result_path):
        raise HTTPException(status_code=404, detail="Hasil file belum tersedia.")

    filename_download = f"{Path(job.filename).stem} - Hasil Parafrase Turnitin.docx"
    return FileResponse(
        path=job.result_path,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename=filename_download
    )

@app.get("/api/jobs/{job_id}/report")
def get_job_report(job_id: str, db: Session = Depends(get_db)):
    job = db.query(ParaphraseJob).filter_by(id=job_id).first()
    if not job:
        raise HTTPException(status_code=404, detail="Job tidak ditemukan.")

    typos = db.query(TypoReport).filter_by(job_id=job.id).all()
    sample_diffs = []

    # Get sample changed paragraphs
    changed = db.query(JobParagraph).filter_by(job_id=job.id, is_locked=False).limit(6).all()
    for c in changed:
        sample_diffs.append({
            "index": c.paragraph_index,
            "original": c.original_text,
            "paraphrased": c.paraphrased_text
        })

    return {
        "job_id": job.id,
        "filename": job.filename,
        "validation_score": job.validation_score,
        "total_paragraphs": job.total_paragraphs,
        "narrative_count": job.narrative_count,
        "locked_count": job.locked_count,
        "typo_findings": [
            {"paragraph": t.paragraph_index, "typo": t.typo_found, "suggestion": t.saran_perbaikan, "note": t.catatan}
            for t in typos
        ],
        "sample_diffs": sample_diffs
    }

# ----------------- DICTIONARY MANAGEMENT ENDPOINTS -----------------
@app.get("/api/dictionary")
def get_dictionary(db: Session = Depends(get_db)):
    padanan = db.query(PadananKata).order_by(PadananKata.kata_asal).all()
    locked = db.query(IstilahTerkunci).order_by(IstilahTerkunci.kategori).all()
    return {
        "padanan": [
            {"id": p.id, "kata_asal": p.kata_asal, "padanan_disukai": p.padanan_disukai, "padanan_terlarang": p.padanan_terlarang, "konteks": p.konteks, "is_active": p.is_active}
            for p in padanan
        ],
        "locked_terms": [
            {"id": l.id, "istilah": l.istilah, "kategori": l.kategori, "deskripsi": l.deskripsi, "is_active": l.is_active}
            for l in locked
        ]
    }

@app.post("/api/dictionary/padanan")
def add_padanan(payload: dict, db: Session = Depends(get_db)):
    kata_asal = payload.get("kata_asal", "").strip().lower()
    padanan_disukai = payload.get("padanan_disukai", "").strip()
    padanan_terlarang = payload.get("padanan_terlarang", "").strip()
    konteks = payload.get("konteks", "Umum").strip()

    if not kata_asal or not padanan_disukai:
        raise HTTPException(status_code=400, detail="Kata asal dan padanan disukai wajib diisi.")

    existing = db.query(PadananKata).filter_by(kata_asal=kata_asal).first()
    if existing:
        existing.padanan_disukai = padanan_disukai
        existing.padanan_terlarang = padanan_terlarang
        existing.konteks = konteks
    else:
        db.add(PadananKata(
            kata_asal=kata_asal,
            padanan_disukai=padanan_disukai,
            padanan_terlarang=padanan_terlarang,
            konteks=konteks
        ))
    db.commit()
    return {"status": "success", "message": "Padanan kata berhasil disimpan."}

@app.delete("/api/dictionary/padanan/{padanan_id}")
def delete_padanan(padanan_id: int, db: Session = Depends(get_db)):
    item = db.query(PadananKata).filter_by(id=padanan_id).first()
    if item:
        db.delete(item)
        db.commit()
    return {"status": "success", "message": "Padanan kata berhasil dihapus."}

@app.post("/api/dictionary/locked")
def add_locked_term(payload: dict, db: Session = Depends(get_db)):
    istilah = payload.get("istilah", "").strip()
    kategori = payload.get("kategori", "Umum").strip()
    deskripsi = payload.get("deskripsi", "").strip()

    if not istilah:
        raise HTTPException(status_code=400, detail="Istilah wajib diisi.")

    existing = db.query(IstilahTerkunci).filter_by(istilah=istilah).first()
    if not existing:
        db.add(IstilahTerkunci(istilah=istilah, kategori=kategori, deskripsi=deskripsi))
        db.commit()
    return {"status": "success", "message": "Istilah terkunci berhasil disimpan."}

@app.delete("/api/dictionary/locked/{term_id}")
def delete_locked_term(term_id: int, db: Session = Depends(get_db)):
    item = db.query(IstilahTerkunci).filter_by(id=term_id).first()
    if item:
        db.delete(item)
        db.commit()
    return {"status": "success", "message": "Istilah terkunci berhasil dihapus."}

if __name__ == "__main__":
    import uvicorn
    uvicorn.run("main:app", host=HOST, port=PORT, reload=True)
