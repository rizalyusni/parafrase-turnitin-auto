import docx

doc = docx.Document('LAPORAN wawancara revisi 3.docx')

print('--- TABLES ---')
for t_idx, t in enumerate(doc.tables):
    print(f'Table {t_idx}: {len(t.rows)} rows, {len(t.columns)} cols')
    for r_idx, r in enumerate(t.rows[:2]):
        row_txt = [c.text.strip().replace('\n', ' ') for c in r.cells]
        print(f'  r{r_idx}: {row_txt[:4]}')

print('\n--- IMAGES / DRAWINGS ---')
img_count = 0
for i, p in enumerate(doc.paragraphs):
    if 'w:drawing' in p._p.xml:
        print(f'Paragraph [{i}] contains drawing/image: {repr(p.text.strip()[:50])}')
        img_count += 1
print(f'Total drawing paragraphs: {img_count}')

print('\n--- SECTIONS & MARGINS ---')
for s_idx, sec in enumerate(doc.sections):
    print(f'Sec {s_idx}: top={sec.top_margin.cm:.2f}cm, bot={sec.bottom_margin.cm:.2f}cm, left={sec.left_margin.cm:.2f}cm, right={sec.right_margin.cm:.2f}cm, w={sec.page_width.cm:.2f}cm, h={sec.page_height.cm:.2f}cm')
