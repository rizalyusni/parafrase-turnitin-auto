# -*- coding: utf-8 -*-
"""
Script to apply paraphrases to LAPORAN wawancara revisi 3.docx
and save to LAPORAN wawancara revisi 3 - Hasil Parafrase.docx
"""
import re
import docx

from para_laporan import LAPORAN_PARAPHRASES

# Foreign/legal Latin terms to italicize
ITALIC_TERMS = [
    r"hoofdelijk voor het geheel",
    r"persona standi in judicio",
    r"limited liability",
    r"Commanditaire Vennootschap",
    r"Online Single Submission",
    r"annual objectives",
    r"annual objective",
    r"job description",
    r"jobdesc",
    r"person in charge",
    r"due process of law",
]

pattern_str = "(" + "|".join(ITALIC_TERMS) + ")"
pattern = re.compile(pattern_str, re.IGNORECASE)

def update_paragraph(paragraph, text):
    """
    Safely updates text in paragraph, preserving style and paragraph formatting,
    applying Times New Roman font and italics to specified terms.
    """
    p_elem = paragraph._p
    
    # Check if there are any runs with drawings
    drawing_runs = []
    for r in list(p_elem.xpath('./w:r')):
        if r.xpath('./w:drawing'):
            drawing_runs.append(r)
        else:
            p_elem.remove(r)
            
    # Add new text runs
    tokens = pattern.split(text)
    for token in tokens:
        if not token:
            continue
        run = paragraph.add_run(token)
        run.font.name = 'Times New Roman'
        if pattern.fullmatch(token.strip()):
            run.italic = True
        else:
            run.italic = False

def main():
    source_filename = "LAPORAN wawancara revisi 3.docx"
    output_filename = "LAPORAN wawancara revisi 3 - Hasil Parafrase.docx"
    
    doc = docx.Document(source_filename)
    print(f"Loaded: {len(doc.paragraphs)} paragraphs, {len(doc.tables)} tables, {len(doc.sections)} sections.")
    
    updated_count = 0
    for idx, new_text in LAPORAN_PARAPHRASES.items():
        if idx < len(doc.paragraphs):
            p = doc.paragraphs[idx]
            update_paragraph(p, new_text)
            updated_count += 1
        else:
            print(f"Warning: Index {idx} out of range!")
            
    print(f"Total paragraphs updated: {updated_count}")
    
    doc.save(output_filename)
    print(f"Saved: {output_filename}")

if __name__ == "__main__":
    main()
