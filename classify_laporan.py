import docx

doc = docx.Document('LAPORAN wawancara revisi 3.docx')

locked_indices = set()

# Front matter & Cover (0 to 11)
locked_indices.update(range(12))

# Kata Pengantar closing & meta:
locked_indices.update([12, 17, 18, 19, 20, 21, 22, 23, 24])

# Daftar Isi (25 to 50)
locked_indices.update(range(25, 51))

# Headings and labels:
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if not txt:
        locked_indices.add(i)
        continue
    if txt in [
        'BAB I', 'PENDAHULUAN', 'LATAR BELAKANG', 'RUMUSAN MASALAH', 'TUJUAN', 'METODE PENELITIAN',
        'BAB II', 'LANDASAN TEORI', 'PERUSAHAAN PERORANGAN (UD)', '2.2 PERSEKUTUAN KOMANDITER (CV)',
        '2.3 PERSEROAN TERBATAS (PT)', '2.4 LEGALITAS DAN PERIZINAN BERUSAHA', '2.5 PERBANDINGAN SINGKAT',
        'BAB III', 'TABEL HASIL WAWANCARA DAN OBSERVASI', 'IDENTITAS PERUSAHAAN', 'LEGALITAS PERUSAHAAN',
        'STRUKTUR ORGANISASI DAN TATA KELOLA', 'PERMASALAHAN HUKUM', 'HASIL OBSERVASI',
        'BAB IV', 'ANALISIS', '4.1 ANALISIS BENTUK HUKUM', '4.2 ANALISIS STRUKTUR PENDIRIAN DAN TANGGUNG JAWAB',
        '4.3 ANALISIS LEGALITAS', '4.4  ANALISIS TATA KELOLA', '4.5 ANALISIS KETENAGAKERJAAN',
        'BAB V', 'PENUTUP', '5.1 KESIMPULAN', '5.2 SARAN', 'DAFTAR PUSTAKA', 'LAMPIRAN'
    ]:
        locked_indices.add(i)
    elif i in [89, 90, 91, 92, 102]: # Interviewee metadata & observation placeholder
        locked_indices.add(i)
    elif i >= 142: # Daftar Pustaka & Lampiran
        locked_indices.add(i)

cand_indices = [i for i in range(len(doc.paragraphs)) if i not in locked_indices]
print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Locked paragraphs: {len(locked_indices)}")
print(f"Candidate paragraphs to paraphrase: {len(cand_indices)}")

with open('laporan_candidates.txt', 'w', encoding='utf-8') as f:
    for idx in cand_indices:
        f.write(f"[{idx}] {repr(doc.paragraphs[idx].text.strip())}\n")

print("Saved laporan_candidates.txt")
