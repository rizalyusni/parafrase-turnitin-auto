import docx

doc = docx.Document('LAPORAN wawancara revisi 3.docx')

with open('dump_laporan.txt', 'w', encoding='utf-8') as f:
    f.write(f'TOTAL PARAGRAPHS: {len(doc.paragraphs)}\n')
    f.write(f'TOTAL TABLES: {len(doc.tables)}\n')
    f.write(f'TOTAL SECTIONS: {len(doc.sections)}\n\n')
    for i, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        f.write(f'[{i}] ({p.style.name}, runs={len(p.runs)}) {txt}\n')

print('Dumped dump_laporan.txt')
