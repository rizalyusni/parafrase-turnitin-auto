import docx

doc = docx.Document('Tesis Bab IV - VI.docx')

print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

with open('dump_tesis_all.txt', 'w', encoding='utf-8') as f:
    for i, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        f.write(f"[{i}] ({p.style.name}, runs={len(p.runs)}) {txt}\n")

print("Dumped dump_tesis_all.txt")
