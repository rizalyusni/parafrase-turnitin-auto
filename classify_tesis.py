import docx
import re

doc = docx.Document('Tesis Bab IV - VI.docx')

locked_indices = set()

# Pattern for formulas or math calculations:
formula_pattern = re.compile(r'^(Rumus|r\s*=|d\s*=|d\s+max|d\s+min|di\s*=|w\s*=|FPIi\s*=)', re.IGNORECASE)

for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if not txt:
        locked_indices.add(i)
        continue
    
    # Headings / Subheadings:
    if p.style.name.startswith('Heading') or txt in [
        'BAB IV', 'HASIL PENELITIAN', 'Deskripsi Objek Penelitian', 'Penyajian Data',
        'Capital Adequacy Ratio (CAR)', 'Return On Asset (ROA)', 'Return On Equity (ROE)',
        'Beban Operasional Terhadap Pendapatan Operasional  (BOPO)',
        'Beban Operasional Terhadap Pendapatan Operasional (BOPO)',
        'Financing to Deposit Ratio (FDR)', 'Zakat Performance Ratio (ZPR)',
        'BAB V', 'ANALISIS DAN PEMBAHASAN', 'Analisis Data', 'Penyusunan Matriks Keputusan',
        'Normalisasi Matriks Keputusan', 'Normalisasi Matriks Keputusan ',
        'Normalisasi Tahun 2021', ' Normalisasi Tahun 2021', 'Normalisasi CAR',
        'Normalisasi Retrun On Asset (ROA)', 'Normalisasi Return On Asset (ROA)',
        'Normalisasi Return On Equity (ROE)', 'Normalisasi BOPO',
        'Normalisasi Financial to Deposit Ratio (FDR) 2021', 'Normalisasi ZPR 2021',
        'Matriks Normalisasi dan Perhitungan FPI Tahun 2021',
        'Normalisasi Matriks Keputusan Tahun 2022', 'Normalisasi CAR 2022',
        'Normalisasi ROA 2022', 'Normalisasi ROE 2022', 'Normalisasi BOPO 2022',
        'Normalisasi FDR 2022', 'Normalisasi ZPR 2022',
        'Matriks Normalisasi dan Perhitungan FPI Tahun 2022',
        'Peringkat Falah Performance Index (FPI) Tahun 2022',
        'Penyusunan Matriks Keputusan 2023', 'Normalisasi CAR 2023',
        'Normalisasi ROA Tahun 2023', 'Normalisasi ROE Tahun 2023',
        'Normalisasi BOPO 2023', 'Normalisasi FDR 2023', 'Normalisasi ZPR 2023',
        'Matriks Normalisasi dan Perhitungan FPI 2023',
        'Peringkat Falah Performance Index (FPI) Tahun 2023',
        'Normalisasi CAR 2024', 'Normalisasi ROA 2024', 'Normalisasi ROE 2024',
        'Normalisasi BOPO 2024', 'Normalisasi FDR 2024', 'Normalisasi ZPR 2024',
        'Matriks Normalisasi dan Perhitungan FPI 2024',
        'Normalisasi CAR 2025', 'Normalisasi ROA 2025', 'Normalisasi ROE 2025',
        'Normalisasi BOPO 2025', 'Normalisasi FDR 2025', 'Normalisasi ZPR 2025',
        'Penyusunan Matriks Keputusan FPI 2025',
        'Pembahasan', 'Pembahasan ',
        'Perkembangan Kinerja Bank Umum Syariah di Indonesia Berdasarkan Indikator CAR, ROA, ROE, BOPO, FDR, dan ZPR Periode 2021 – 2025',
        'Perkembangan Kinerja Bank Umum Syariah di Indonesia Berdasarkan Indikator CAR, ROA, ROE, BOPO, FDR, dan ZPR Periode 2021  2025',
        'Perkembangan Capital Adequacy Ratio (CAR)',
        'Perkembangan Return On Assets (ROA)',
        'Perkembangan Return On Equity (ROA)',
        'Perkembangan Return On Equity (ROE)',
        'Perkembangan Biaya Operasional terhadap Pendapatan Operasional (BOPO)',
        'Perkembangan Financing to Deposit Ratio (FDR)',
        'Perkembangan Zakat Performance Ratio (ZPR)',
        'Perkembangan Kinerja Bank Umum Syariah Berdasarkan 6 Indikator',
        'Hasil Pengukuran Kinerja Bank Umum Syariah Melalui Falah Performance Index (FPI) Periode 2021 – 2025',
        'Hasil Pengukuran Kinerja Bank Umum Syariah Melalui Falah Performance Index (FPI) Periode 2021  2025',
        'Peringkat Kinerja Bank Umum Syariah Berdasarkan Falah Performance Index (FPI) Periode 2021 – 2025',
        'Peringkat Kinerja Bank Umum Syariah Berdasarkan Falah Performance Index (FPI) Periode 2021  2025',
        'Peringkat Tahun 2021', 'Peringkat Tahun 2022', 'Peringkat Tahun 2023',
        'Peringkat Tahun 2024', 'Peringkat Tahun 2025',
        'Dinamika Peringkat Selama Periode 2021–2025',
        'Dinamika Peringkat Selama Periode 20212025',
        'Pencapaian Prinsip Falah Berdasarkan Hasil Pengukuran Falah Performance Index (FPI) pada Bank Umum Syariah Periode 2021–2025',
        'Pencapaian Prinsip Falah Berdasarkan Hasil Pengukuran Falah Performance Index (FPI) pada Bank Umum Syariah Periode 20212025',
        'Falah dalam Aspek Ketahanan dan Keberlanjutan Keuangan',
        'Falah dalam Aspek Produktivitas dan Profitabilitas',
        'Falah dalam Aspek Efisiensi Pengelolaan Sumber Daya',
        'Falah dalam Fungsi Intermediasi dan Kemaslahatan Ekonomi',
        'Falah dalam Dimensi Sosial melalui Zakat',
        'Keseimbangan Kinerja sebagai Cerminan Falah',
        'Makna Perbedaan Nilai FPI terhadap Pencapaian Falah',
        'Sintesis Pencapaian Prinsip Falah'
    ]:
        locked_indices.add(i)
        continue
    
    # Table titles and sources
    if txt.startswith('Tabel ') or txt.startswith('Tabel\t') or txt == 'Tabel' or txt.startswith('Sumber :') or txt.startswith('Sumber:'):
        locked_indices.add(i)
        continue
    if txt.startswith('Data Tabel') or txt.startswith('Data Matriks') or txt.startswith('Data Normalisasi') or txt.startswith('Tabel Normalisasi') or txt.startswith('Tabel Data Normalisasi') or txt.startswith('Matriks Normalisasi') or txt.startswith('Perkembangan Peringkat'):
        locked_indices.add(i)
        continue
    if txt.startswith('Peringkat Falah Performance Index') or txt.startswith('Peringkat FPI'):
        locked_indices.add(i)
        continue

    # Formulas / math calculations
    if formula_pattern.match(txt) or (len(txt) < 30 and any(c in txt for c in ['=', '<', '>', '≤', '≥', '–', '-'])):
        locked_indices.add(i)
        continue

print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Locked paragraphs: {len(locked_indices)}")
cand_indices = [i for i in range(len(doc.paragraphs)) if i not in locked_indices]
print(f"Candidate paragraphs to paraphrase: {len(cand_indices)}")

with open('tesis_candidates.txt', 'w', encoding='utf-8') as f:
    for idx in cand_indices:
        p = doc.paragraphs[idx]
        f.write(f"[{idx}] ({p.style.name}) {p.text.strip()}\n")

print("Saved tesis_candidates.txt")
